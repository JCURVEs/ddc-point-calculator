from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parent))

from ddc_selector import DdcInput, select_ddc, write_excel_output


class DdcSelectorTests(unittest.TestCase):
    def test_40_points_uses_pxc4(self) -> None:
        selected = select_ddc(DdcInput(ddc_no="DDC-S", di=10, do=10, ai=10, ao=10), "TXM1.8U")

        self.assertEqual(selected.source.total, 40)
        self.assertEqual(selected.controller, "PXC4.E16S-2")

    def test_47_points_uses_pxc5(self) -> None:
        selected = select_ddc(DdcInput(ddc_no="DDC-1", di=16, do=8, ai=16, ao=7), "TXM1.8X")

        self.assertEqual(selected.source.total, 47)
        self.assertEqual(selected.controller, "PXC5.E24")
        self.assertEqual(selected.txm16d_qty, 1)
        self.assertEqual(selected.txm6r_qty, 2)
        self.assertEqual(selected.analog_module, "TXM1.8X")
        self.assertEqual(selected.analog_module_qty, 3)

    def test_84_points_exceeds_pxc5_and_uses_pxc7s(self) -> None:
        selected = select_ddc(DdcInput(ddc_no="DDC-9", di=28, do=14, ai=30, ao=12), "TXM1.8X")

        self.assertEqual(selected.source.total, 84)
        self.assertEqual(selected.controller, "PXC7.E400S")
        self.assertEqual(selected.txm16d_qty, 2)
        self.assertEqual(selected.txm6r_qty, 3)
        self.assertEqual(selected.analog_module_qty, 6)

    def test_controller_override_is_supported(self) -> None:
        selected = select_ddc(
            DdcInput(ddc_no="DDC-X", di=10, do=10, ai=10, ao=10, controller_override="PXC7.E400M"),
            "TXM1.8U",
        )

        self.assertEqual(selected.controller, "PXC7.E400M")

    def test_excel_output_contains_no_material_sheet(self) -> None:
        from openpyxl import load_workbook

        selected = select_ddc(DdcInput(ddc_no="DDC-1", di=16, do=8, ai=16, ao=7), "TXM1.8X")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "selection.xlsx"
            write_excel_output(path, [selected])
            workbook = load_workbook(path, read_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["DDC Selection", "Rules"])
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
